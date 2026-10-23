from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import logging

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_PATH = BASE_DIR / "output" / "reporte_cuentas.xlsx"


def generar_excel(registros):
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte de cuentas"

    encabezados = [
        "linea_csv",
        "valores_recibidos",
        "valores_limpios",
        "estado",
        "userId",
        "title",
        "body",
        "causa",
    ]
    hoja.append(encabezados)

    for encabezado in hoja[1]:
        encabezado.font = Font(bold=True, color="FFFFFF")
        encabezado.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        encabezado.alignment = Alignment(horizontal="center")

    for registro in registros:
        datos = registro.get("datos") or {}
        hoja.append([
            registro["linea"],
            registro["valores_recibidos"],
            registro["valores_limpios"],
            registro["estado"],
            datos.get("userId"),
            datos.get("title"),
            datos.get("body"),
            registro.get("causa"),
        ])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    anchos = {"A": 10, "B": 18, "C": 14, "D": 18, "E": 10, "F": 40, "G": 70, "H": 24}
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    try:
        libro.save(OUTPUT_PATH)
    except OSError as error:
        logging.exception("No se pudo guardar el reporte Excel.")
        raise RuntimeError(f"No se pudo guardar el reporte Excel en {OUTPUT_PATH}: {error}") from error

    print(f"\nReporte Excel generado: {OUTPUT_PATH}")
    logging.info("Reporte Excel generado correctamente: %s", OUTPUT_PATH)